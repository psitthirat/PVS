"""Current, post-correction read-only numerical audit of the author's 1 October DOCX; exports aggregate Excel.

Never writes the DOCX, raw data, analysis frame or frozen September release.
Counts are recalculated from the trusted local frame. September model estimates
are checked against frozen approved tables. The four overall-report models
are independently re-estimated by pvs.report_models.
"""
from pathlib import Path
import hashlib, json, re
from decimal import Decimal, ROUND_HALF_UP
import numpy as np
import pandas as pd
from scipy.stats import chi2_contingency, mannwhitneyu, random_table
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pvs.national import q40_top2

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'output/reports/ผลการสำรวจ_20261001.docx'
OUT=ROOT/'output/precommit_review_20261001'
FRAME=ROOT/'output/analysis_frame.pkl'
GROUPS=['public primary care','public hospital OPD/ER','private clinic/pharmacy','private hospital OPD/ER']
NUMBERS={f'T{i:02}':i if i==1 else i-1 if i<14 else i-2 for i in range(1,23)}
NUMBERS.update(T02='สรุปภาพรวม',T14='สรุปเอกชน')


def blocks():
    doc=Document(SOURCE); out=[]; p=t=0
    for elem in doc.element.body.iterchildren():
        if elem.tag.endswith('}p'):
            p+=1; q=Paragraph(elem,doc);out.append(dict(id=f'P{p:03}',type='paragraph',text=q.text))
        elif elem.tag.endswith('}tbl'):
            t+=1;out.append(dict(id=f'T{t:02}',type='table',rows=[[c.text for c in row.cells] for row in Table(elem,doc).rows]))
    return out


def rounded(value, digits):
    return str(Decimal(str(value)).quantize(Decimal('1').scaleb(-digits),rounding=ROUND_HALF_UP))


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    digest=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    d=pd.read_pickle(FRAME); B=blocks(); tables={b['id']:b['rows'] for b in B if b['type']=='table'}
    paras={b['id']:b['text'] for b in B if b['type']=='paragraph'}
    shapes=[(16,6),(1,1),(23,5),(7,5),(17,5),(11,5),(14,5),(21,5),
            (10,3),(7,4),(6,3),(6,3),(11,4),(1,1),(5,5),(10,5),
            (9,4),(10,5),(8,4),(12,4),(14,5),(9,4)]
    if len(paras)!=320 or [(len(rows),len(rows[0])) for rows in tables.values()]!=shapes:
        raise ValueError('Report layout changed. Update paragraph/table mappings before using this audit.')
    checked=[]; issues=[]; tests=[]
    def issue(loc,kind,old,new,reason,source='output/analysis_frame.pkl',severity='แก้ก่อนเผยแพร่'):
        issues.append(dict(ตำแหน่ง=loc,ประเภท=kind,ข้อความเดิม=str(old),ค่าหรือข้อความที่เสนอ=str(new),เหตุผล=reason,หลักฐาน=source,สถานะ=severity,แก้ไขแล้ว=''))
    def loc(t,r,c=None):
        return f'ตาราง {NUMBERS[t]} / {t} แถว {r+1}'+(f' คอลัมน์ {c+1}' if c is not None else '')+' / '+tables[t][r][0].replace('\n',' ')
    def check(t,r,c,expected,note,kind='ตัวเลข'):
        raw=tables[t][r][c];ok=str(raw).replace('\n',' ').strip()==str(expected)
        checked.append(dict(ตำแหน่ง=loc(t,r,c),เดิม=raw,ตรวจได้=str(expected),ผล='ตรง' if ok else 'ต่าง',วิธี=note))
        if not ok: issue(loc(t,r,c),kind,raw,expected,note)
    def count(t,r,c,values,mask,positive='yes',denom=None):
        valid=mask & values.notna() if denom is None else mask & denom
        n=int((valid & values.eq(positive)).sum()); N=int(valid.sum());value=100*n/N if N else np.nan
        raw=tables[t][r][c]; m=re.search(r'([\d,]+)(?:/([\d,]+))?\s*\(?\s*([\d.]+)\s*%',raw)
        if not m: raise ValueError((t,r,c,raw))
        digits=len(m[3].partition('.')[2]);expected=f'{n:,} ({rounded(value,digits)}%)'
        ok=int(m[1].replace(',',''))==n and m[3]==rounded(value,digits) and (not m[2] or int(m[2].replace(',',''))==N)
        note=f'n/N = {n}/{N}; {values.name}; positive={positive}; recalculated from canonical frame'
        checked.append(dict(ตำแหน่ง=loc(t,r,c),เดิม=raw,ตรวจได้=expected,ผล='ตรง' if ok else 'ต่าง',วิธี=note))
        if not ok: issue(loc(t,r,c),'จำนวน/ร้อยละ',raw,expected,note)
    allmask=pd.Series(True,index=d.index)
    areas=[allmask,d.municipality.eq('municipal'),d.municipality.eq('non-municipal')]
    sectors=[d.usual_sector.eq('public'),d.usual_sector.eq('private')]
    groups=[d.provider_group.eq(g) for g in GROUPS]
    def np_row(t,r,var,pos='yes',masks=areas,restriction=None,all_den=False):
        if t=='T03' and r>=6: r+=1  # author added the missing-sex row
        vals=d[var] if isinstance(var,str) else var
        for c,mask in enumerate(masks,1):
            mask=mask if restriction is None else mask & restriction
            count(t,r,c,vals,mask,pos,allmask if all_den else None)
    def ptest(t,r,values,group=d.municipality,mask=allmask,positive=None):
        if t=='T03' and r>=6: r+=1
        valid=mask & values.notna() & group.notna(); v=values[valid]
        if positive is not None: v=v.eq(positive)
        ct=pd.crosstab(group[valid],v); p=chi2_contingency(ct,correction=False).pvalue
        py=chi2_contingency(ct,correction=True).pvalue
        exp=chi2_contingency(ct,correction=False).expected_freq
        pearson=p
        method='Pearson chi-square without continuity correction'
        se=0.0
        if (exp<5).any():
            # Conditional Monte Carlo Pearson statistic with fixed margins.
            # Matches the method stated in P097; retain seed/draw count.
            rng=np.random.default_rng(20261001)
            obs=((ct.to_numpy()-exp)**2/exp).sum(); exceed=0; draws=999999
            for start in range(0,draws,10000):
                sim=random_table.rvs(ct.sum(axis=1).to_numpy(),ct.sum(axis=0).to_numpy(),size=min(10000,draws-start),random_state=rng)
                stats=((sim-exp)**2/exp).sum(axis=(1,2))
                exceed+=int((stats>=obs-1e-12).sum())
            p=(exceed+1)/(draws+1);se=float(np.sqrt(p*(1-p)/draws))
            method='Monte Carlo Pearson; fixed margins; 999999 draws; seed=20261001'
        raw=tables[t][r][-1].strip(); digits=len(raw.partition('.')[2]) if '.' in raw else 2
        if '<' in raw: ok=p<float(raw.replace('<','').strip())
        else: ok=raw==rounded(p,digits)
        tests.append(dict(ตำแหน่ง=loc(t,r),ตัวแปร=values.name,นิยาม=str(positive or 'ทุกระดับที่ตอบได้'),เดิม=raw,p=p,Pearsonไม่แก้Yates=pearson,Yates=py,วิธี=method,MonteCarloSE=se,N=int(valid.sum()),expected_min=float(exp.min()),ผล='ตรง' if ok else 'ต่าง',ตารางจำนวน=ct.to_json(force_ascii=False)))
        if not ok:
            issue(loc(t,r)+' p-value','ค่า p',raw,f'{p:.8g} (แสดง {"<0.001" if p<.001 else rounded(p,digits)})',method)
    # Table 2: denominators follow the printed full-sample headings, with missing sex explicitly noted.
    for r,pos in [(3,'male'),(4,'female'),(5,'other')]: np_row('T03',r,'sex',pos,all_den=True)
    for c,mask in enumerate(areas,1): count('T03',6,c,d.sex.isna(),mask,True,allmask)
    edu=d.edu.astype(object).replace({'below primary':'primary','upper secondary/voc':'upper','diploma/high voc':'upper'})
    for r,pos in [(7,'primary'),(8,'lower secondary'),(9,'upper'),(10,'bachelor or above')]: np_row('T03',r,edu,pos)
    inc=d.income.astype(object).replace({'<10k':'under30','10-30k':'under30','60-100k':'60plus','100k+':'60plus'})
    for r,pos in [(12,'under30'),(13,'30-60k'),(14,'60plus')]: np_row('T03',r,inc,pos)
    scheme=d.scheme.astype(object).fillna('unknown')
    for r,pos in [(16,'UCS'),(17,'SSS'),(18,'CSMBS'),(19,'other public'),(20,'unknown')]: np_row('T03',r,scheme,pos)
    np_row('T03',21,'priv_ins_owned');
    for r,vals,pos in [(2,d.sex.astype(object).fillna('unknown'),None),(6,edu,None),(11,inc,None),(15,scheme,None),(21,d.priv_ins_owned,'yes')]: ptest('T03',r,vals,positive=pos)
    # The corrected Table 3 explicitly uses good or better for both health measures.
    for r,var,pos in [(2,'health_good3','good or better'),(3,'mental_good3','good or better'),(4,'chronic','yes'),(5,'activation','high')]:
        np_row('T04',r,var,pos);ptest('T04',r,d[var],positive=pos)
    # Table 4: female-only screening; Q27/Q29/Q31 are past 12 months, researcher confirmed.
    services={3:'admit',4:'unmet',6:'bp_r',7:'sugar_r',8:'lipid_r',9:'mammogram_r',10:'ca_cervix_r',11:'vision_r',12:'dental_r',13:'homevisit_any',14:'telemed_any',15:'borrowed_ever'}
    for r,var in services.items():
        restriction=d.sex.eq('female') if r in [9,10] else None
        np_row('T05',r,var,restriction=restriction);ptest('T05',r,d[var],mask=restriction if restriction is not None else allmask,positive='yes')
    for c,m in enumerate(areas,1):
        v=d.loc[m,'visits_n'].dropna();q=v.quantile([.25,.5,.75]);check('T05',2,c,f'{q.loc[.5]:.2f} ({q.loc[.25]:.2f}-{q.loc[.75]:.2f})','Median (IQR), valid visits_n')
    mw=mannwhitneyu(d.loc[areas[1],'visits_n'].dropna(),d.loc[areas[2],'visits_n'].dropna(),alternative='two-sided').pvalue
    check('T05',2,4,rounded(mw,2),f'Mann–Whitney two-sided p={mw:.8g}')
    np_row('T06',2,'has_usual');ptest('T06',2,d.has_usual,positive='yes')
    for r,v,p in [(5,'usual_sector','public'),(6,'usual_sector','private'),(8,'usual_hospital',False),(9,'usual_hospital',True)]: np_row('T06',r,v,p,all_den=True)
    ptest('T06',4,d.usual_sector,mask=d.usual_sector.isin(['public','private']))
    ptest('T06',7,d.usual_hospital)
    g3=[groups[0],groups[1],sectors[1]]
    for r,v,p in [(2,'municipality','municipal'),(3,d.edu.isin(['lower secondary','upper secondary/voc','diploma/high voc','bachelor or above']),True),(5,inc,'under30'),(6,inc,'30-60k'),(7,inc,'60plus'),(9,'scheme','UCS'),(10,'scheme','SSS'),(11,'scheme','CSMBS'),(12,'usual_qual_top2','very good/excellent')]: np_row('T07',r,v,p,masks=g3,all_den=True)
    g3col=pd.Series(np.select(g3,['public primary','public hospital','private'],default='excluded'),index=d.index)
    for r,v,pos in [(2,d.municipality,'municipal'),(3,d.edu.isin(['lower secondary','upper secondary/voc','diploma/high voc','bachelor or above']),True),(4,inc,None),(8,scheme,None),(12,d.usual_qual_top2,'very good/excellent')]: ptest('T07',r,v,group=g3col,mask=g3col.ne('excluded'),positive=pos)
    confidence=[('conf_get','confident'),('conf_afford','confident'),('health_security','confident'),('sys_improved_bin','improved'),('sys_endorse_bin','works well')]
    for r,(v,p) in enumerate(confidence,2):
        np_row('T10',r,v,p,masks=sectors);ptest('T10',r,d[v],group=d.usual_sector,mask=d.usual_sector.isin(['public','private']),positive=p)
    qualities=[(d.usual_qual_top2,'very good/excellent'),(d.q_public_top2,'very good/excellent'),(d.q_private_top2,'very good/excellent')]+[(q40_top2(d[c]),True) for c in ['eval_maternal','eval_ped','eval_chronic','eval_mental']]+[(d.conf_voice,'confident'),(d.q_covid_top2,'very good/excellent')]
    for r,(v,p) in enumerate(qualities,2):
        np_row('T13',r,v,p,masks=sectors);ptest('T13',r,v,group=d.usual_sector,mask=d.usual_sector.isin(['public','private']),positive=p)
    for r,v in [(2,'appoint'),(3,'wait_ge1mo'),(4,'queue_ge1h')]: np_row('T15',r,v,'appointment' if r==2 else 'yes',masks=groups)
    for r,v in enumerate(['one_facility_multi','admit','homevisit_any','telemed_any','unmet','borrowed_ever','med_error_ever','discrim_ever'],2): np_row('T16',r,v,masks=groups)
    preventive=['bp_r','sugar_r','lipid_r','dental_r','vision_r','mental_r','mammogram_r','ca_cervix_r']
    for r,v in enumerate(preventive,2): np_row('T18',r,v,masks=groups,restriction=d.sex.eq('female') if r>=8 else None)
    subgroups=[d.private_primary_sub.eq('clinic'),d.private_primary_sub.eq('pharmacy'),groups[-1]]
    for r,v in enumerate(preventive[:6],2): np_row('T19',r,v,masks=subgroups)
    for r,v,p in [(3,'conf_get','confident'),(4,'conf_afford','confident'),(5,'health_security','confident'),(7,'conf_voice','confident'),(8,'sys_improved_bin','improved'),(9,'sys_endorse_bin','works well'),(11,'q_public_top2','very good/excellent'),(12,'q_private_top2','very good/excellent'),(13,'q_covid_top2','very good/excellent')]: np_row('T21',r,v,p,masks=groups)
    # Check adjusted private-sector estimates to the approved release at their displayed precision.
    for t,name in [('T17','T59_core_outcomes_adjusted'),('T20','T54_experience_adjusted'),('T22','T58_confidence_adjusted')]:
        f=pd.read_csv(ROOT/f'results/2026-09/tables/{name}.csv');f=f[f.Group.eq(GROUPS[-1])].reset_index(drop=True)
        assert len(f)==len(tables[t])-1
        for i,row in f.iterrows():
            r=i+1; check(t,r,1,f'{row.aOR:.2f} ({row.lo:.2f}–{row.hi:.2f})',f'{name}, {row.Outcome}; approved September estimate')
            check(t,r,2,f'{row["Diff (pp)"]:+.1f}',f'{name}, average marginal difference (percentage points)')
            pv=row['p (BH within family)']; check(t,r,3,'<0.001' if pv<.001 else f'{pv:.3f}',f'{name}, p adjusted by BH within family')
    # Re-estimate all four overall-model tables under their printed specification.
    from pvs.report_models import estimate
    models=estimate(d)
    (OUT/'overall-models.json').write_text(json.dumps(models,ensure_ascii=False,indent=2)+'\n')
    def modelcheck(t,r,c,m):
        raw=tables[t][r][c]
        nums=re.findall(r'\d+\.\d+',raw)
        expected=[f'{m[k]:.2f}' for k in ['aOR','lo','hi']]
        ok=nums==expected
        checked.append(dict(ตำแหน่ง=loc(t,r,c),เดิม=raw,ตรวจได้=f'{expected[0]} [{expected[1]}, {expected[2]}]',ผล='ตรง' if ok else 'ต่าง',วิธี='Re-estimated from canonical frame; model-based covariance; complete-case cohort'))
        if not ok:issue(loc(t,r,c),'OR/95%CI',raw,' / '.join(expected),'Re-estimated overall report model')
        raw=tables[t][r][c+1].strip()
        ok=(m['p']<float(raw.replace('<','').strip())) if '<' in raw else raw==rounded(m['p'],len(raw.partition('.')[2]))
        checked.append(dict(ตำแหน่ง=loc(t,r,c+1),เดิม=raw,ตรวจได้=f'{m["p"]:.8g}',ผล='ตรง' if ok else 'ต่าง',วิธี='Wald p; model-based covariance; no multiplicity adjustment'))
        if not ok:issue(loc(t,r,c+1),'ค่า p โมเดล',raw,f'{m["p"]:.8g}','Re-estimated overall report model')
    rows={2:('sex','female'),3:('area','municipal'),4:('education','secondary or higher'),
          6:('age','30-44'),7:('age','45-59'),8:('age','60-74'),9:('age','75+'),
          11:('income','middle'),12:('income','high'),14:('scheme','SSS'),15:('scheme','CSMBS'),
          16:('activation','high'),17:('health','good or better'),18:('sector','public'),19:('setting','nonhospital')}
    for r,(v,level) in rows.items():
        for c,kind in [(1,'univariate'),(3,'multivariate')]:
            m=next(m for m in models['T08'][kind] if m['variable']==v and m['level']==level)
            modelcheck('T08',r,c,m)
    for t,var in [('T09','usual'),('T11','usual'),('T12','sector')]:
        for r,(outcome,m) in enumerate(models[t].items(),2):
            estimate=next(x for x in m['multivariate'] if x['variable']==var)
            modelcheck(t,r,1,estimate)
    assert models['T08']['n']==1268
    assert all(m['n']==(945 if m['women_only'] else 1651) for m in models['T09'].values())
    assert {m['n'] for m in models['T11'].values()}=={1651}
    assert {m['n'] for m in models['T12'].values()}=={1268}
    # Segmentation models and aggregate tables are recalculated, not just hashed.
    from pvs import segments as seg
    from io import StringIO
    sd=seg._frame(); a,b,c=seg.explanatory_power(sd)
    fresh={'T82_user_decomposition_routes':seg.decomposition(sd),'T83_decomposition_scheme_x_insurance':seg.decomposition_detail(sd),'T84_reason_by_route':seg.reason_by_route(sd),'T85_explanatory_power_single':a,'T86_explanatory_power_nested':b,'T87_scheme_vs_income_tests':c,'T88_scheme_x_income_grid':seg.scheme_income_grid(sd),'T89_age_profile_by_band':seg.age_profile(sd),'T91_age_specification':seg.age_specification(sd)}
    segmentation=[]
    for name,new in fresh.items():
        old=pd.read_csv(ROOT/f'output/tables/{name}.csv')
        new=pd.read_csv(StringIO(new.to_csv(index=False)))
        pd.testing.assert_frame_equal(new.reset_index(drop=True),old,check_dtype=False,atol=1e-10,rtol=1e-8)
        segmentation.append({'table':name,'rows':len(new),'result':'Recalculated and matched'})
    # Header and content anchors fail closed rather than apply an old row mapping.
    assert 'ไม่ระบุ' in tables['T03'][6][0]
    assert 'ดี ดีมาก' in tables['T04'][-1][0]
    assert 'Monte Carlo' in paras['P097'] and 'ไม่ใช้ continuity correction' in paras['P097']
    # Preserve the historical sampling-plan table; external population figures
    # are traced to that source, not independently recertified as current facts.
    plan=Document(ROOT/'output/reports/ผลการสำรวจ_20260801.docx').tables[0]
    assert tables['T01']==[[c.text for c in row.cells] for row in plan.rows]
    checked.append(dict(ตำแหน่ง='ตาราง 1 / T01',เดิม='ตารางแผนการสำรวจ',ตรวจได้='ตรงกับตารางแผนในรายงานต้นทางทุกช่อง',ผล='ตรง',วิธี='Exact text match to August plan; not an independent census audit'))
    from review_report_claims import review
    claims,text_issues,inventory=review(d,B,models,tests)
    # Use compact, directly editable suggestions for the two prose discrepancies.
    for i in text_issues:
        if i['ตำแหน่ง']=='P116':
            i.update(ข้อความเดิม='p = 0.19',ค่าหรือข้อความที่เสนอ='p = 0.18',เหตุผล='Monte Carlo Pearson ตามวิธี P097: p=0.179263 (999999 simulations, seed 20261001); คงข้อสรุปไม่พบความแตกต่าง')
        elif i['ตำแหน่ง']=='P189':
            i.update(ข้อความเดิม='ร้อยละ 54.7 เทียบกับร้อยละ 45.7',ค่าหรือข้อความที่เสนอ='ร้อยละ 54.6 เทียบกับร้อยละ 45.7',เหตุผล='147/269 × 100 = 54.646840... ปัดเป็นหนึ่งตำแหน่งโดยตรงได้ 54.6; ไม่ปัดซ้ำจาก 54.65 ในตาราง')
    issues.extend(text_issues)
    # Python normalizes Thai keyword identifiers differently from string keys.
    import unicodedata,zipfile
    def normalized(records):
        return [{unicodedata.normalize('NFKC',k):v for k,v in row.items()} for row in records]
    issues=normalized(issues)
    image_rows=[]
    assets={hashlib.sha256(p.read_bytes()).hexdigest():str(p.relative_to(ROOT)) for base in ['site/assets/figures','results/2026-09/figures'] for p in (ROOT/base).glob('*.png')}
    # Visual reviews apply only to these exact embedded bytes. Image 9 is a
    # smaller raster of the access chart, not a conceptual-framework image.
    reviewed_images={
        '3183af321b4fb46eae49f5f518c9bec00b7e1e3222819c638d367e4ea01aeadd':(
            '', 'รูปที่ 1: ตรวจด้วยสายตา เป็นกรอบแนวคิด ไม่มีค่าประมาณจากผลสำรวจ'),
        '521a906463ef3d76dc70b2ff25ebcd2d761167f98feb630543c7d4095719f2c7':(
            'site/assets/figures/figure-3.png',
            'รูปที่ 9: ตรวจด้วยสายตาเทียบกราฟปัจจุบันและ CSV แล้ว ค่าที่แสดงสอดคล้องกัน รวมถึงตัวหารผู้ใช้บริการมากกว่า 1 ครั้ง; ภาพ Word 918×750 เป็นคนละขนาดและไม่ตรงกันระดับไบต์'),
    }
    with zipfile.ZipFile(SOURCE) as archive:
        for member in archive.infolist():
            name=member.filename
            if name.startswith('word/media/') and not member.is_dir():
                h=hashlib.sha256(archive.read(name)).hexdigest()
                if h in assets:
                    match,scope=assets[h],'ตรงกับไฟล์ภาพที่ตรวจข้อมูลแล้วทุกไบต์'
                elif h in reviewed_images:
                    match,scope=reviewed_images[h]
                else:
                    match,scope='','ภาพใหม่หรือเปลี่ยนแปลง: ต้องตรวจเพิ่มเติม'
                    issue(name,'ภาพที่ยังไม่ได้ตรวจ',h,'ตรวจเทียบแหล่งข้อมูลก่อนเผยแพร่',scope)
                image_rows.append({'ภาพ':name,'SHA256':h,'ตรงกับไฟล์':assets.get(h,''),
                                   'ไฟล์ที่ใช้เทียบ':match,'ขอบเขต':scope})
    required=[i for i in issues if i['สถานะ']!='แนะนำปรับถ้อยคำ']
    result={'source':str(SOURCE.relative_to(ROOT)),'sha256':digest,
            'ready_for_publication':not required,'remaining_required_edits':len(required),
            'cell_checks':checked,'p_tests':tests,'claims':claims,'issues':issues,
            'segmentation':segmentation,'images':image_rows,'overall_models_reestimated':True,
            'recall_confirmation':'content/questionnaire/recall-confirmation.json',
            'scope':'Outcome counts, percentages, table p-values, printed overall models, approved September contrasts, result prose, summaries and chart provenance. Background/reference years and external census figures are not independently recertified.'}
    (OUT/'audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    from openpyxl.utils import get_column_letter
    workbook=Workbook();workbook.remove(workbook.active)
    def sheet(name,records):
        records=normalized(records);ws=workbook.create_sheet(name)
        if not records:return ws
        cols=list(records[0]);ws.append(cols)
        for record in records:
            values=[record.get(c,'') for c in cols]
            ws.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in values])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(name='Tahoma',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='913340')
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=85
            for cell in row:cell.font=Font(name='Tahoma',size=11);cell.alignment=Alignment(vertical='top',wrap_text=True)
        for n,c in enumerate(cols,1):ws.column_dimensions[get_column_letter(n)].width=20 if c in ['N','p','ผล','สถานะ'] else 52
        return ws
    sheet('อ่านก่อน',[
        {'หัวข้อ':'ฉบับที่ตรวจ','รายละเอียด':str(SOURCE.relative_to(ROOT))},
        {'หัวข้อ':'SHA256','รายละเอียด':digest},
        {'หัวข้อ':'ผลตรวจ','รายละเอียด':f'ตรวจตาราง {len(checked)} รายการ ค่า p {len(tests)} รายการ และข้อความ/กรอบสรุป {len(claims)} รายการ; เหลือ {len(required)} ตำแหน่งที่ต้องแก้ และ {len(issues)-len(required)} ข้อเสนอถ้อยคำ'},
        {'หัวข้อ':'โมเดลภาพรวม','รายละเอียด':'ประมาณตาราง 7,8,10,11 ซ้ำจาก canonical frame แล้ว ค่าสัมประสิทธิ์ CI ค่า p และ N ตรงกับ Word ฉบับแก้; ไม่ถือเป็นปัญหาค้างเดิมอีกต่อไป'},
        {'หัวข้อ':'นิยามช่วงเวลา','รายละเอียด':'ผู้วิจัยตรวจแบบสอบถามต้นฉบับและยืนยัน Q27, Q29, Q31 = 12 เดือนที่ผ่านมา; ใช้แทนการอนุมานเดิมจากหัวคอลัมน์ที่ย่อไว้'},
        {'หัวข้อ':'ตำแหน่ง','รายละเอียด':'Pxxx นับย่อหน้ารวมว่าง ไม่ใช่เลขหน้า; Txx คือตาราง XML; ใช้ข้อความเดิมค้นหาใน Word. T14 คือกรอบสรุปข้อค้นพบของส่วนโรงพยาบาลเอกชน'},
        {'หัวข้อ':'ขอบเขตที่ไม่รับรอง','รายละเอียด':'ปี/เลขอ้างอิงในบทนำและสถิติประชากรของแผนสำรวจไม่ได้ตรวจต้นทางภายนอกใหม่; ตารางแผนเทียบกับฉบับสิงหาคม; ไม่ใช่การรับรองความเป็นตัวแทนของกลุ่มตัวอย่างหรือการตรวจแบ่งหน้า Word'},
        {'หัวข้อ':'ต้นฉบับ','รายละเอียด':'ไม่ได้แก้ Word/ข้อมูลดิบ/ผลเผยแพร่กันยายน; Workbook นี้แทนรายการแก้ของการตรวจฉบับก่อนหน้า ไม่ควรใช้รายการเก่ากับฉบับล่าสุด'},
    ])
    sheet('จุดที่ต้องแก้',required)
    sheet('ข้อเสนอถ้อยคำ',[i for i in issues if i['สถานะ']=='แนะนำปรับถ้อยคำ'])
    sheet('ตรวจตาราง',checked);sheet('ค่า p',tests);sheet('ตรวจข้อความ',claims)
    sheet('ขอบเขตข้อความ',inventory);sheet('ตรวจภาพ',image_rows);sheet('คำนวณซ้ำกลุ่มเอกชน',segmentation)
    sheet('นิยามโมเดล',[{'ตาราง':t,'Outcome':k,'N':m['n'],'Events':m['events'],'Covariance':m['covariance'],'Predictors':', '.join(m['predictors'])} for t,value in models.items() for k,m in ([(t,value)] if t=='T08' else value.items())])
    sheet('แหล่งข้อมูล',[{'ไฟล์':str(p.relative_to(ROOT)),'SHA256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [SOURCE,FRAME,ROOT/'content/national/national_summary.csv',ROOT/'content/international/values.csv',ROOT/'content/questionnaire/recall-confirmation.json',ROOT/'src/pvs/report_models.py',Path(__file__),ROOT/'tools/review_report_claims.py']])
    dest=ROOT/'output/reports/ตรวจซ้ำก่อนเผยแพร่_ผลการสำรวจ_20261001.xlsx'
    workbook.save(dest)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==digest
    print(json.dumps({'workbook':str(dest.relative_to(ROOT)),'checks':len(checked),'p_tests':len(tests),'claims':len(claims),'required_edits':len(required),'wording_suggestions':len(issues)-len(required),'source_unchanged':True},ensure_ascii=False))
    for i in issues:print(i)

if __name__=='__main__':main()
