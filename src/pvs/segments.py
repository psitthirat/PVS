"""
Who is the main target population of private hospitals?

    python -m pvs segments      (needs the tables stage first)

Outcome   : usual source of care is a private hospital OPD/ER (n = 127).
Base      : ALL 2,017 respondents, so a rate answers "how likely is someone in this
            group to be a regular private-hospital user". People with no usual source
            of care count as non-users, which is what a population target requires.

Every cell reports two different things that answer two different questions:
  * Rate  (penetration)  users / everyone in the cell      -> who is most LIKELY
  * Share (composition)  users in cell / all 127 users     -> where the users ARE
  * Index                rate / overall rate (6.3%) x 100  -> 200 = twice as likely
A good target is high on both rate and share.
"""
from __future__ import annotations
import os, itertools, warnings
import numpy as np, pandas as pd
import statsmodels.formula.api as smf
from .analysis import prop_ci, bh

warnings.filterwarnings('ignore')
from .paths import OUTPUT as OUT
TAB = f'{OUT}/tables'
SMALL_N = 30          # cells with fewer people than this are flagged as unstable

DIMS = {
    'Age':            ('age_u50', {'yes': 'Under 50', 'no': '50 and over'}),
    'Area':           ('municipality', {'municipal': 'Municipal (urban)',
                                        'non-municipal': 'Non-municipal (rural)'}),
    'Scheme':         ('scheme_gr', {'SSS': 'Social Security', 'UCS': 'Universal Coverage',
                                     'CSMBS': 'Civil servant', "other/don't know": "Other / don't know"}),
    'Income':         ('income3', {'<10k': 'Under 10k', '10-30k': '10k-30k', '30k+': '30k and over'}),
    'Education':      ('edu_gr', {'upper secondary or above': 'Upper secondary +',
                                  'lower secondary or below': 'Lower secondary or less'}),
    'Sex':            ('sex', {'female': 'Female', 'male': 'Male', 'other': 'Other'}),
    'Private insurance': ('priv_ins_owned', {'yes': 'Owns private insurance',
                                             'no': 'No private insurance'}),
}
PAIRS = [('Age', 'Area'), ('Age', 'Scheme'), ('Area', 'Scheme'), ('Income', 'Scheme'),
         ('Age', 'Income'), ('Education', 'Scheme'), ('Private insurance', 'Scheme'),
         ('Area', 'Income'), ('Sex', 'Age')]


def _frame():
    d = pd.read_pickle(f'{OUT}/analysis_frame.pkl').copy()
    d['user'] = (d['provider_group'] == 'private hospital OPD/ER').astype(int)
    for name, (col, lab) in DIMS.items():
        d[name] = d[col].map(lab)
    return d


def _cell(sub, total_users, overall):
    n, k = len(sub), int(sub['user'].sum())
    rate, lo, hi = prop_ci(k, n) if n else (np.nan, np.nan, np.nan)
    return {'People (N)': n, 'Private-hospital users (n)': k,
            'Rate %': round(rate, 1), 'Rate 95% CI low': round(lo, 1), 'Rate 95% CI high': round(hi, 1),
            'Share of all users %': round(100 * k / total_users, 1),
            'Index (100 = average)': round(100 * (rate / 100) / overall) if n else np.nan,
            'Small cell (N < 30)': 'yes' if n < SMALL_N else ''}


def one_way(d):
    tot, overall = int(d['user'].sum()), d['user'].mean()
    rows = []
    for name in DIMS:
        for level, sub in d.dropna(subset=[name]).groupby(name, sort=False):
            rows.append({'Characteristic': name, 'Group': level, **_cell(sub, tot, overall)})
    return pd.DataFrame(rows)


def two_way(d, a, b):
    """Long table plus three presentation matrices (rate, share, count) for one pair."""
    tot, overall = int(d['user'].sum()), d['user'].mean()
    dd = d.dropna(subset=[a, b])
    la = [v for v in DIMS[a][1].values() if v in dd[a].unique()]
    lb = [v for v in DIMS[b][1].values() if v in dd[b].unique()]
    rows = []
    for x in la:
        for y in lb:
            sub = dd[(dd[a] == x) & (dd[b] == y)]
            rows.append({a: x, b: y, **_cell(sub, tot, overall)})
    long = pd.DataFrame(rows)

    def matrix(value_fn, total_fn):
        m = pd.DataFrame(index=la + ['All'], columns=lb + ['All'], dtype=object)
        for x in la + ['All']:
            for y in lb + ['All']:
                sub = dd
                if x != 'All':
                    sub = sub[sub[a] == x]
                if y != 'All':
                    sub = sub[sub[b] == y]
                m.loc[x, y] = value_fn(sub)
        m.index.name = f'{a} \\ {b}'
        return m.reset_index()

    def rate_txt(sub):
        n, k = len(sub), int(sub['user'].sum())
        if n == 0:
            return '--'
        flag = ' *' if n < SMALL_N else ''
        return f'{100 * k / n:.1f}%{flag}'

    rate_m = matrix(rate_txt, None)
    share_m = matrix(lambda s: f"{100 * s['user'].sum() / tot:.1f}%", None)
    count_m = matrix(lambda s: f"{int(s['user'].sum())} / {len(s)}", None)
    return long, rate_m, share_m, count_m


def segments(d, dims):
    """Every combination of `dims`, ranked by rate, with cumulative reach."""
    tot, overall = int(d['user'].sum()), d['user'].mean()
    dd = d.dropna(subset=dims)
    rows = []
    for combo, sub in dd.groupby(dims, sort=False):
        combo = combo if isinstance(combo, tuple) else (combo,)
        rows.append({'Segment': ' + '.join(combo), **dict(zip(dims, combo)),
                     **_cell(sub, tot, overall)})
    s = pd.DataFrame(rows).sort_values(['Rate %', 'Private-hospital users (n)'],
                                       ascending=[False, False]).reset_index(drop=True)
    s.insert(0, 'Rank', range(1, len(s) + 1))
    s['Cumulative share of users %'] = s['Share of all users %'].cumsum().round(1)
    s['Cumulative share of population %'] = (100 * s['People (N)'].cumsum() / len(dd)).round(1)
    return s


def target_definitions(d):
    """Candidate target definitions compared on precision (rate) and reach (share)."""
    tot, overall = int(d['user'].sum()), d['user'].mean()
    sss = d['scheme_gr'] == 'SSS'
    u50 = d['age_u50'] == 'yes'
    urb = d['municipality'] == 'municipal'
    inc = d['income3'] == '30k+'
    pin = d['priv_ins_owned'] == 'yes'
    cands = [
        ('Everyone (reference)', pd.Series(True, index=d.index)),
        ('Under 50', u50),
        ('Municipal (urban)', urb),
        ('Under 50 + municipal', u50 & urb),
        ('Income 30k and over', inc),
        ('Owns private insurance', pin),
        ('Social Security', sss),
        ('Social Security + under 50', sss & u50),
        ('Social Security + municipal', sss & urb),
        ('Social Security + under 50 + municipal  (the starting hypothesis)', sss & u50 & urb),
        ('Social Security + income 30k and over', sss & inc),
        ('Social Security OR owns private insurance', sss | pin),
    ]
    rows = []
    for lab, mask in cands:
        sub = d[mask.fillna(False)]
        c = _cell(sub, tot, overall)
        rows.append({'Target definition': lab, 'People (N)': c['People (N)'],
                     '% of population': round(100 * c['People (N)'] / len(d), 1),
                     'Private-hospital users (n)': c['Private-hospital users (n)'],
                     'Rate % (precision)': c['Rate %'],
                     'Rate 95% CI low': c['Rate 95% CI low'], 'Rate 95% CI high': c['Rate 95% CI high'],
                     'Share of all users % (reach)': c['Share of all users %'],
                     'Index (100 = average)': c['Index (100 = average)']})
    return pd.DataFrame(rows)


def decomposition(d):
    """
    An exhaustive (MECE) account of every private-hospital user.

    A single "best" segment can only ever describe a minority of users. These three
    routes partition the whole population, so the user shares sum to 100% and nobody
    is left unexplained. The split is entitlement first, then whether the respondent
    holds private insurance, because those are the two markers that survive adjustment.
    """
    tot = int(d['user'].sum())
    sss = d['scheme_gr'].eq('SSS')
    pins = d['priv_ins_owned'].eq('yes')
    routes = [
        ('A. Social Security (entitlement route)', sss,
         'Scheme registration with a contracted hospital is the plausible mechanism, '
         'but registration is not measured'),
        ('B. Other scheme + holds private insurance', ~sss & pins,
         'Private cover alongside a public entitlement; ownership is measured, use is not'),
        ('C. Other scheme, no private insurance', ~sss & ~pins,
         'No private cover recorded; choice is the likely route, payment unknown'),
    ]
    rows = []
    for lab, mask, note in routes:
        sub = d[mask.fillna(False)]
        k, n = int(sub['user'].sum()), len(sub)
        rate, lo, hi = prop_ci(k, n)
        rows.append({'Route': lab, 'People (N)': n, '% of population': round(100 * n / len(d), 1),
                     'Private-hospital users (n)': k,
                     'Share of all users %': round(100 * k / tot, 1),
                     'Rate % (within route)': round(rate, 1),
                     'Rate 95% CI low': round(lo, 1), 'Rate 95% CI high': round(hi, 1),
                     'Mechanism suggested (not established)': note})
    out = pd.DataFrame(rows)
    total = {'Route': 'All routes (exhaustive)', 'People (N)': len(d), '% of population': 100.0,
             'Private-hospital users (n)': tot, 'Share of all users %': 100.0,
             'Rate % (within route)': round(100 * tot / len(d), 1),
             'Rate 95% CI low': np.nan, 'Rate 95% CI high': np.nan,
             'Mechanism suggested (not established)': ''}
    return pd.concat([out, pd.DataFrame([total])], ignore_index=True)


def decomposition_detail(d):
    """The same partition opened out by scheme and private insurance; sums to 100%."""
    tot = int(d['user'].sum())
    rows = []
    for scheme in ['SSS', 'UCS', 'CSMBS', "other/don't know"]:
        for ins in ['yes', 'no']:
            sub = d[d['scheme_gr'].eq(scheme) & d['priv_ins_owned'].eq(ins)]
            if len(sub) == 0:
                continue
            k, n = int(sub['user'].sum()), len(sub)
            rate, lo, hi = prop_ci(k, n)
            rows.append({'Scheme': scheme,
                         'Private insurance': 'owns' if ins == 'yes' else 'none',
                         'People (N)': n, 'Private-hospital users (n)': k,
                         'Share of all users %': round(100 * k / tot, 1),
                         'Rate %': round(rate, 1), 'Rate 95% CI low': round(lo, 1),
                         'Rate 95% CI high': round(hi, 1),
                         'Small cell (N < 30)': 'yes' if n < SMALL_N else ''})
    out = pd.DataFrame(rows).sort_values('Share of all users %', ascending=False)
    out['Cumulative share of users %'] = out['Share of all users %'].cumsum().round(1)
    return out.reset_index(drop=True)


def reason_by_route(d):
    """Why each route says it chose the facility. Different routes, different reasons."""
    u = d[d['user'] == 1].copy()
    u['Route'] = np.where(u['scheme_gr'].eq('SSS'), 'A. Social Security',
                          np.where(u['priv_ins_owned'].eq('yes'),
                                   'B. Other scheme + private insurance',
                                   'C. Other scheme, no private insurance'))
    ct = pd.crosstab(u['Route'], u['usual_reason'])
    pct = (100 * ct.div(ct.sum(axis=1), axis=0)).round(1)
    out = []
    for route in ct.index:
        for reason in ct.columns:
            if ct.loc[route, reason] == 0:
                continue
            out.append({'Route': route, 'Stated reason for choosing the facility': reason,
                        'Users (n)': int(ct.loc[route, reason]),
                        '% of that route': pct.loc[route, reason]})
    return pd.DataFrame(out).sort_values(['Route', '% of that route'], ascending=[True, False])


SOCIO = ['age_u50', 'municipality', 'income3', 'edu_gr', 'sex']


def explanatory_power(d):
    """
    Can characteristics other than the coverage scheme account for private-hospital use?

    Single variables are compared on McFadden pseudo-R-squared, and nested models show
    what the scheme adds to socio-demographics and vice versa. This answers directly
    whether residence, income or education could stand in for entitlement.
    """
    from scipy import stats as st
    labels = {'scheme_gr': 'Coverage scheme', 'priv_ins_owned': 'Owns private insurance',
              'income3': 'Household income', 'municipality': 'Area of residence',
              'age_u50': 'Age under 50', 'edu_gr': 'Education', 'sex': 'Sex'}
    single = []
    tot = int(d['user'].sum())
    for col, lab in labels.items():
        dd = d.dropna(subset=[col])
        m = smf.logit(f'user ~ C({col})', data=dd).fit(disp=0)
        g = dd.groupby(col)['user'].agg(['sum', 'count'])
        g['rate'] = 100 * g['sum'] / g['count']
        top = g.sort_values('rate', ascending=False).iloc[0]
        single.append({'Characteristic': lab, 'Groups': len(g),
                       'McFadden pseudo-R2': round(m.prsquared, 4),
                       'Lowest group rate %': round(g['rate'].min(), 1),
                       'Highest group rate %': round(g['rate'].max(), 1),
                       'Highest group: share of all users %': round(100 * top['sum'] / tot, 1),
                       'n': int(m.nobs)})
    single = pd.DataFrame(single).sort_values('McFadden pseudo-R2', ascending=False)

    dd = d.dropna(subset=SOCIO + ['scheme_gr', 'priv_ins_owned'])
    def fit(cols):
        return smf.logit('user ~ ' + ' + '.join(f'C({c})' for c in cols), data=dd).fit(disp=0)
    models = {
        'Socio-demographics only (age, area, income, education, sex)': fit(SOCIO),
        'Socio-demographics + private insurance': fit(SOCIO + ['priv_ins_owned']),
        'Socio-demographics + coverage scheme': fit(SOCIO + ['scheme_gr']),
        'Socio-demographics + scheme + private insurance': fit(SOCIO + ['scheme_gr', 'priv_ins_owned']),
    }
    nested = [{'Model': k, 'McFadden pseudo-R2': round(m.prsquared, 4),
               'AIC': round(m.aic, 1), 'n': int(m.nobs)} for k, m in models.items()]

    mi, ms = fit(['income3']), fit(['scheme_gr'])
    mb = fit(['income3', 'scheme_gr'])
    tests = []
    for lab, small in [('Adding coverage scheme to an income-only model', mi),
                       ('Adding income to a scheme-only model', ms)]:
        lr = 2 * (mb.llf - small.llf)
        df = mb.df_model - small.df_model
        tests.append({'Test': lab, 'Likelihood-ratio chi2': round(lr, 1),
                      'df': int(df), 'p': f'{st.chi2.sf(lr, df):.2g}'})
    return single, pd.DataFrame(nested), pd.DataFrame(tests)


def scheme_income_grid(d):
    """Rate by coverage scheme and household income: two gradients, not one."""
    rows = []
    for inc in ['<10k', '10-30k', '30k+']:
        for sch in ['SSS', 'UCS', 'CSMBS']:
            sub = d[d['income3'].eq(inc) & d['scheme_gr'].eq(sch)]
            if len(sub) == 0:
                continue
            k, n = int(sub['user'].sum()), len(sub)
            rate, lo, hi = prop_ci(k, n)
            rows.append({'Household income': inc, 'Coverage scheme': sch,
                         'People (N)': n, 'Private-hospital users (n)': k,
                         'Rate %': round(rate, 1), 'Rate 95% CI low': round(lo, 1),
                         'Rate 95% CI high': round(hi, 1),
                         'Small cell (N < 30)': 'yes' if n < SMALL_N else ''})
    return pd.DataFrame(rows)


def continuity_multi_visit(d):
    """
    Use of a single facility, among respondents with MORE THAN ONE visit.

    Q20 ("were all your visits to the same facility?") was fielded to everyone with at
    least one visit. Respondents who attended exactly once answer "one facility" as a
    matter of arithmetic, which inflates every group. Restricting to more than one
    visit states the denominator plainly and removes that artefact.
    """
    from scipy import stats as st
    coh = d[d['provider_group'].notna()]
    rows = []
    for lab, sub in [('All who answered Q20 (one or more visits)', coh[coh['one_facility'].notna()]),
                     ('More than one visit (analysis definition)', coh[coh['visits_n'] > 1])]:
        for g in prep_groups():
            s_ = sub.loc[sub['provider_group'] == g, 'one_facility'].dropna()
            k, n = int((s_ == 'yes').sum()), len(s_)
            rate, lo, hi = prop_ci(k, n)
            rows.append({'Denominator': lab, 'Provider group': g,
                         'Respondents in denominator (N)': n, 'Used one facility (n)': k,
                         '% using one facility': round(rate, 1),
                         '95% CI low': round(lo, 1), '95% CI high': round(hi, 1)})
        ct = pd.crosstab(sub['one_facility'], sub['provider_group'])
        rows.append({'Denominator': lab, 'Provider group': 'All groups (chi-square)',
                     'Respondents in denominator (N)': int(ct.values.sum()),
                     'Used one facility (n)': int(ct.loc['yes'].sum()),
                     '% using one facility': round(100 * ct.loc['yes'].sum() / ct.values.sum(), 1),
                     '95% CI low': np.nan, '95% CI high': np.nan,
                     'p': f'{st.chi2_contingency(ct)[1]:.2g}'})
    single = coh[coh['visits_n'] == 1]['one_facility'].dropna()
    rows.append({'Denominator': 'Memo: respondents with exactly one visit',
                 'Provider group': 'All groups',
                 'Respondents in denominator (N)': len(single),
                 'Used one facility (n)': int((single == 'yes').sum()),
                 '% using one facility': round(100 * (single == 'yes').mean(), 1),
                 '95% CI low': np.nan, '95% CI high': np.nan,
                 'p': 'excluded from the analysis definition'})
    return pd.DataFrame(rows)


def prep_groups():
    from .preparation import GROUPS
    return GROUPS


def age_profile(d):
    """
    Age without imposing a cut point.

    The earlier <50 / 50+ split was a convenience, not a finding. Using the full
    seven-band Q4 variable (observed for every respondent) shows the shape directly,
    overall and within the two schemes that hold almost all users.
    """
    rows = []
    frames = [('All respondents', d),
              ('Social Security', d[d['scheme_gr'].eq('SSS')]),
              ('Universal Coverage', d[d['scheme_gr'].eq('UCS')])]
    for lab, frame in frames:
        for band in ['18-29', '30-39', '40-49', '50-59', '60-69', '70-79', '80+']:
            sub = frame[frame['age_band'].astype('object').eq(band)]
            if len(sub) == 0:
                continue
            k, n = int(sub['user'].sum()), len(sub)
            rate, lo, hi = prop_ci(k, n)
            rows.append({'Population': lab, 'Age band': band, 'People (N)': n,
                         'Private-hospital users (n)': k, 'Rate %': round(rate, 1),
                         'Rate 95% CI low': round(lo, 1), 'Rate 95% CI high': round(hi, 1),
                         'Small cell (N < 30)': 'yes' if n < SMALL_N else ''})
    return pd.DataFrame(rows)


def age_specification(d):
    """Is age linear? Does it add anything once the coverage scheme is known?"""
    from scipy import stats as st
    rows = []
    dd = d.dropna(subset=['age_exact', 'scheme_gr', 'income3']).copy()
    dd['age10'] = dd['age_exact'] / 10
    for lab, formula in [('Age alone (continuous, per 10 years)', 'user ~ age10'),
                         ('Age + coverage scheme', 'user ~ age10 + C(scheme_gr)'),
                         ('Age + scheme + income', 'user ~ age10 + C(scheme_gr) + C(income3)')]:
        m = smf.logit(formula, data=dd).fit(disp=0)
        b, se = m.params['age10'], m.bse['age10']
        rows.append({'Specification': lab, 'Odds ratio per 10 years': round(np.exp(b), 2),
                     '95% CI low': round(np.exp(b - 1.959964 * se), 2),
                     '95% CI high': round(np.exp(b + 1.959964 * se), 2),
                     'p': round(float(m.pvalues['age10']), 3),
                     'McFadden pseudo-R2': round(m.prsquared, 4), 'n': int(m.nobs),
                     'Note': 'exact age (Q3), missing for 308 respondents'})
    lin = smf.logit('user ~ age10 + C(scheme_gr)', data=dd).fit(disp=0)
    quad = smf.logit('user ~ age10 + I(age10**2) + C(scheme_gr)', data=dd).fit(disp=0)
    lr = 2 * (quad.llf - lin.llf)
    rows.append({'Specification': 'Test: is the age effect linear? (adding a quadratic term)',
                 'Odds ratio per 10 years': np.nan, '95% CI low': np.nan, '95% CI high': np.nan,
                 'p': round(float(st.chi2.sf(lr, 1)), 3), 'McFadden pseudo-R2': round(quad.prsquared, 4),
                 'n': int(quad.nobs),
                 'Note': f'LR chi2 = {lr:.1f}, df = 1. A significant result means a straight line '
                         'misdescribes the pattern.'})
    db = d.dropna(subset=['age_band', 'scheme_gr'])
    a = smf.logit('user ~ C(scheme_gr)', data=db).fit(disp=0)
    b2 = smf.logit('user ~ C(scheme_gr) + C(age_band)', data=db).fit(disp=0)
    lr2 = 2 * (b2.llf - a.llf)
    df2 = b2.df_model - a.df_model
    rows.append({'Specification': 'Test: does the seven-band age variable add anything to scheme?',
                 'Odds ratio per 10 years': np.nan, '95% CI low': np.nan, '95% CI high': np.nan,
                 'p': round(float(st.chi2.sf(lr2, df2)), 3), 'McFadden pseudo-R2': round(b2.prsquared, 4),
                 'n': int(b2.nobs),
                 'Note': f'LR chi2 = {lr2:.1f}, df = {df2:.0f}. Uses the complete Q4 band, '
                         'so no respondent is dropped.'})
    return pd.DataFrame(rows)


def adjusted(d):
    """Which characteristics matter independently of each other?"""
    dd = d.dropna(subset=['age_u50', 'municipality', 'scheme_gr', 'income3', 'edu_gr',
                          'sex', 'priv_ins_owned', 'subdist']).copy()
    dd = dd[dd['scheme_gr'] != "other/don't know"]
    # "other" sex (n=47) has zero private-hospital users, which separates the model and
    # yields a meaningless odds ratio of 0. Those respondents stay in every descriptive
    # table; only the adjusted model is restricted to male/female.
    dd = dd[dd['sex'].isin(['male', 'female'])]
    f = ('user ~ C(age_u50, Treatment("no")) + C(municipality, Treatment("non-municipal"))'
         ' + C(scheme_gr, Treatment("UCS")) + C(income3, Treatment("<10k"))'
         ' + C(edu_gr, Treatment("lower secondary or below")) + C(sex, Treatment("male"))'
         ' + C(priv_ins_owned, Treatment("no"))')
    model = smf.logit(f, data=dd)
    m = model.fit(disp=0)
    mc = model.fit(disp=0, cov_type='cluster',
                   cov_kwds={'groups': dd['subdist'].astype('category').cat.codes})
    import re
    LAB = {('age_u50', 'yes'): ('Under 50', 'vs 50 and over'),
           ('municipality', 'municipal'): ('Municipal (urban)', 'vs non-municipal'),
           ('scheme_gr', 'SSS'): ('Social Security', 'vs Universal Coverage'),
           ('scheme_gr', 'CSMBS'): ('Civil servant', 'vs Universal Coverage'),
           ('income3', '10-30k'): ('Income 10k-30k', 'vs under 10k'),
           ('income3', '30k+'): ('Income 30k and over', 'vs under 10k'),
           ('edu_gr', 'upper secondary or above'): ('Upper secondary +', 'vs lower secondary or less'),
           ('sex', 'female'): ('Female', 'vs male'),
           ('priv_ins_owned', 'yes'): ('Owns private insurance', 'vs none')}
    REF = {'age_u50': 'no', 'municipality': 'non-municipal', 'scheme_gr': 'UCS',
           'income3': '<10k', 'edu_gr': 'lower secondary or below', 'sex': 'male',
           'priv_ins_owned': 'no'}
    rows = []
    for i, nm in enumerate(m.params.index):
        mm = re.match(r'C\((\w+),.*\)\[T\.(.+)\]$', nm)
        if not mm or (mm.group(1), mm.group(2)) not in LAB:
            continue
        var, lev = mm.group(1), mm.group(2)
        b, se = float(np.asarray(mc.params)[i]), float(np.asarray(mc.bse)[i])
        on, off = dd.copy(), dd.copy()
        on[var], off[var] = lev, REF[var]
        p_on, p_off = float(m.predict(on).mean()), float(m.predict(off).mean())
        lab = LAB[(var, lev)]
        rows.append({'Characteristic': lab[0], 'Compared with': lab[1],
                     'Adjusted odds ratio': round(np.exp(b), 2),
                     '95% CI low': round(np.exp(b - 1.959964 * se), 2),
                     '95% CI high': round(np.exp(b + 1.959964 * se), 2),
                     'p': float(np.asarray(mc.pvalues)[i]),
                     'Predicted rate if yes %': round(100 * p_on, 1),
                     'Predicted rate if reference %': round(100 * p_off, 1),
                     'Difference (pp)': round(100 * (p_on - p_off), 1)})
    r = pd.DataFrame(rows)
    r['p (BH)'] = bh(r['p'].values)
    r['p'] = r['p'].round(4); r['p (BH)'] = r['p (BH)'].round(4)
    return r, int(m.nobs)


def main():
    d = _frame()
    tot, N = int(d['user'].sum()), len(d)
    print(f'{N} respondents, {tot} regular private-hospital users ({100*tot/N:.1f}%)\n')
    res = {}

    ow = one_way(d)
    ow.to_csv(f'{TAB}/T74_target_one_way.csv', index=False, encoding='utf-8-sig')
    res['one_way'] = ow

    xls = {}
    long_all = []
    for a, b in PAIRS:
        long, rate_m, share_m, count_m = two_way(d, a, b)
        long.insert(0, 'Pair', f'{a} x {b}')
        long_all.append(long.rename(columns={a: 'Row', b: 'Column'}))
        xls[f'{a} x {b}'] = (rate_m, share_m, count_m)
    two = pd.concat(long_all, ignore_index=True)
    two.to_csv(f'{TAB}/T75_target_two_way.csv', index=False, encoding='utf-8-sig')
    res['two_way'] = two

    core = segments(d, ['Age', 'Area', 'Scheme'])
    core.to_csv(f'{TAB}/T76_target_segments_age_area_scheme.csv', index=False, encoding='utf-8-sig')
    ext = segments(d, ['Age', 'Area', 'Scheme', 'Income'])
    ext.to_csv(f'{TAB}/T77_target_segments_with_income.csv', index=False, encoding='utf-8-sig')
    tdef = target_definitions(d)
    tdef.to_csv(f'{TAB}/T79_target_definitions_compared.csv', index=False, encoding='utf-8-sig')
    res['tdef'] = tdef
    dec = decomposition(d)
    dec.to_csv(f'{TAB}/T82_user_decomposition_routes.csv', index=False, encoding='utf-8-sig')
    det = decomposition_detail(d)
    det.to_csv(f'{TAB}/T83_decomposition_scheme_x_insurance.csv', index=False, encoding='utf-8-sig')
    rea = reason_by_route(d)
    rea.to_csv(f'{TAB}/T84_reason_by_route.csv', index=False, encoding='utf-8-sig')
    res.update(dec=dec, det=det, rea=rea)
    cont = continuity_multi_visit(d)
    cont.to_csv(f'{TAB}/T92_one_facility_multi_visit.csv', index=False, encoding='utf-8-sig')
    res['cont'] = cont
    ageprof = age_profile(d)
    ageprof.to_csv(f'{TAB}/T89_age_profile_by_band.csv', index=False, encoding='utf-8-sig')
    agespec = age_specification(d)
    agespec.to_csv(f'{TAB}/T91_age_specification.csv', index=False, encoding='utf-8-sig')
    res.update(ageprof=ageprof, agespec=agespec)
    single, nested, tests = explanatory_power(d)
    single.to_csv(f'{TAB}/T85_explanatory_power_single.csv', index=False, encoding='utf-8-sig')
    nested.to_csv(f'{TAB}/T86_explanatory_power_nested.csv', index=False, encoding='utf-8-sig')
    tests.to_csv(f'{TAB}/T87_scheme_vs_income_tests.csv', index=False, encoding='utf-8-sig')
    grid = scheme_income_grid(d)
    grid.to_csv(f'{TAB}/T88_scheme_x_income_grid.csv', index=False, encoding='utf-8-sig')
    res.update(single=single, nested=nested, tests=tests, grid=grid)
    adj, n_adj = adjusted(d)
    adj.to_csv(f'{TAB}/T78_target_adjusted.csv', index=False, encoding='utf-8-sig')
    res.update(core=core, ext=ext, adj=adj, n_adj=n_adj, xls=xls, N=N, tot=tot)
    _workbook(res)
    print('wrote T74-T79, T82-T89, T91-T92 and "Target segmentation - private hospital users.xlsx"')
    return res


def _workbook(res):
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    path = f'{OUT}/Target segmentation - private hospital users.xlsx'
    overall = 100 * res['tot'] / res['N']

    def shade(v):
        """White at the average rate, deepening orange as the rate rises."""
        try:
            x = float(str(v).split('%')[0])
        except ValueError:
            return None
        t = max(0.0, min(1.0, (x - overall) / 30))
        r, g, b = 255, int(255 - t * (255 - 168)), int(255 - t * (255 - 64))
        if x < overall:
            u = max(0.0, min(1.0, (overall - x) / overall))
            r, g, b = int(255 - u * 30), int(255 - u * 20), 255
        return PatternFill('solid', fgColor=f'{r:02X}{g:02X}{b:02X}')

    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        # ---- README sheet
        intro = pd.DataFrame({'How to read this workbook': [
            f'Base: all {res["N"]:,} respondents. A private-hospital user is someone whose usual source '
            f'of care is a private hospital OPD/ER (n = {res["tot"]}); overall rate {overall:.1f}%.',
            'RATE = users / everyone in the cell. Answers: who is most LIKELY to be a user?',
            'SHARE = users in the cell / all users. Answers: where are the users?',
            'INDEX = rate / overall rate x 100. 200 means twice as likely as average.',
            'A good target is high on BOTH rate and share. A tiny group can have a high rate but few users.',
            '* marks cells with fewer than 30 people: the rate is unstable, read it as indicative only.',
            'Shading on the rate matrices: orange = above the overall rate, blue = below.',
            'Unweighted quota-based convenience sample; these describe the achieved sample, not Thailand.',
            'Scheme is coverage OWNERSHIP. It does not show registration with, or payment to, the hospital.',
        ]})
        intro.to_excel(xw, sheet_name='READ ME', index=False)
        xw.sheets['READ ME'].column_dimensions['A'].width = 120

        res['cont'].to_excel(xw, sheet_name='One facility (>1 visit)', index=False)
        res['ageprof'].to_excel(xw, sheet_name='Age profile by band', index=False)
        res['agespec'].to_excel(xw, sheet_name='Age specification', index=False)
        res['single'].to_excel(xw, sheet_name='Explanatory power', index=False)
        res['nested'].to_excel(xw, sheet_name='Nested models', index=False)
        res['tests'].to_excel(xw, sheet_name='Scheme vs income', index=False)
        res['grid'].to_excel(xw, sheet_name='Scheme x income grid', index=False)
        res['dec'].to_excel(xw, sheet_name='Decomposition (100%)', index=False)
        res['det'].to_excel(xw, sheet_name='Decomposition detail', index=False)
        res['rea'].to_excel(xw, sheet_name='Reason by route', index=False)
        res['tdef'].to_excel(xw, sheet_name='Target definitions', index=False)
        res['core'].to_excel(xw, sheet_name='Segments age-area-scheme', index=False)
        res['ext'].to_excel(xw, sheet_name='Segments + income', index=False)
        res['adj'].to_excel(xw, sheet_name='Adjusted (independent effects)', index=False)
        res['one_way'].to_excel(xw, sheet_name='One-way', index=False)

        for pair, (rate_m, share_m, count_m) in res['xls'].items():
            sheet = pair[:31]
            r0 = 0
            for title, m, shade_it in [('RATE — % of people in the cell who are private-hospital users', rate_m, True),
                                       ('SHARE — % of all private-hospital users who fall in the cell', share_m, False),
                                       ('COUNT — users / people in the cell', count_m, False)]:
                pd.DataFrame({title: []}).to_excel(xw, sheet_name=sheet, startrow=r0, index=False)
                m.to_excel(xw, sheet_name=sheet, startrow=r0 + 1, index=False)
                ws = xw.sheets[sheet]
                ws.cell(row=r0 + 1, column=1).font = Font(bold=True, size=11)
                if shade_it:
                    for i in range(len(m)):
                        for j in range(1, m.shape[1]):
                            c = ws.cell(row=r0 + 3 + i, column=j + 1)
                            f = shade(c.value)
                            if f is not None:
                                c.fill = f
                r0 += len(m) + 4
            ws.column_dimensions['A'].width = 30
            for j in range(2, 8):
                ws.column_dimensions[get_column_letter(j)].width = 21

        for name in ['One facility (>1 visit)', 'Age profile by band', 'Age specification', 'Explanatory power', 'Nested models', 'Scheme vs income',
                     'Scheme x income grid', 'Decomposition (100%)', 'Decomposition detail', 'Reason by route',
                     'Target definitions', 'Segments age-area-scheme', 'Segments + income', 'Adjusted (independent effects)', 'One-way']:
            ws = xw.sheets[name]
            ws.freeze_panes = 'B2'
            for j in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(j)].width = 17
            ws.column_dimensions['B'].width = 46


if __name__ == '__main__':
    main()
